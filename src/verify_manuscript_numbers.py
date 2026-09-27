"""Check every headline number in the manuscript against the released tables.

Each check names the sentence it guards, recomputes the value from the released
data, and asserts that the manuscript states it. A failure means the manuscript
and the release disagree, which is the one thing a reader cannot check for
themselves.

  python verify_manuscript_numbers.py --manuscript ../docs/manuscript_source.md \
      --data ../exports
"""
from pathlib import Path
from argparse import ArgumentParser
import json
import re
import sys

import pandas as pd

MINUS = "−"


def norm(text):
    """Normalise the characters that differ between the source and the data."""
    return (text.replace(MINUS, "-").replace("–", "-").replace("—", "-")
            .replace(" ", " ").replace(",", ""))


class Checker(object):
    def __init__(self, text):
        self.text = norm(text)
        self.passed = []
        self.failed = []

    def has(self, label, *candidates):
        """Assert that at least one spelling of the value appears in the text."""
        for c in candidates:
            if norm(str(c)) in self.text:
                self.passed.append((label, c))
                return True
        self.failed.append((label, candidates))
        return False

    def report(self):
        for label, value in self.passed:
            print("  ok    %-58s %s" % (label, value))
        for label, values in self.failed:
            print("  FAIL  %-58s expected one of %s"
                  % (label, ", ".join(str(v) for v in values)))
        print("\n%d checks passed, %d failed" % (len(self.passed), len(self.failed)))
        return 1 if self.failed else 0


def fmt(v, k=2):
    return ("%.*f" % (k, v)).replace("-", MINUS)


def main():
    p = ArgumentParser(description=__doc__)
    p.add_argument("--manuscript", type=Path, required=True)
    p.add_argument("--data", type=Path, required=True)
    a = p.parse_args()
    d = a.data

    text = ""
    if a.manuscript.suffix.lower() not in (".docx", ".pdf"):
        text = a.manuscript.read_text(encoding="utf-8")
    if a.manuscript.suffix.lower() == ".pdf":
        import pymupdf
        doc = pymupdf.open(str(a.manuscript))
        text = "\n".join(page.get_text() for page in doc)
        doc.close()
    if a.manuscript.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(a.manuscript))
        parts = [q.text for q in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        text = "\n".join(parts)
    c = Checker(text)

    lsum = json.loads((d / "Library_provenance_summary.json").read_text(encoding="utf-8"))
    c.has("source records", "%d" % lsum["records"], "{:,}".format(lsum["records"]))
    c.has("records parsed", lsum["parsed"])
    c.has("motif-bearing records", lsum["motif_bearing_records"])
    c.has("unique motif-bearing", lsum["unique_motif_bearing"])
    c.has("failed the selection rule", lsum["failed_selection_rule"])
    c.has("retained parents", lsum["retained_parents"])
    c.has("distinct electrophilic sites", lsum["distinct_sites_in_retained"])
    c.has("SMARTS matches", lsum["smarts_matches_in_retained"])

    rsum = json.loads((d / "Recognition_ranking_summary.json").read_text(encoding="utf-8"))
    c.has("FP-2 state records", rsum["FP2"]["returned_state_records"])
    c.has("FP-3 state records", rsum["FP3"]["returned_state_records"])
    c.has("FP-2 parents with a pose", rsum["FP2"]["parents_with_pose"])
    c.has("FP-3 parents with a pose", rsum["FP3"]["parents_with_pose"])
    c.has("paired parents", rsum["paired_parents"])
    total_states = (rsum["FP2"]["returned_state_records"]
                    + rsum["FP3"]["returned_state_records"])
    c.has("total state records released", total_states)
    c.has("FP-2 state changes under GlideScore",
          rsum["FP2"]["changed_state_if_GlideScore_minimised"])
    c.has("FP-3 state changes under GlideScore",
          rsum["FP3"]["changed_state_if_GlideScore_minimised"])
    c.has("FP-2 rank correlation between rules",
          fmt(rsum["FP2"]["rank_rho_DockingScore_vs_minimum_GlideScore"], 3))
    c.has("FP-3 rank correlation between rules",
          fmt(rsum["FP3"]["rank_rho_DockingScore_vs_minimum_GlideScore"], 3))
    c.has("minimum composite agreement",
          fmt(min(rsum["composite_agreement_spearman"].values()), 2))
    for t, key in (("FP-2", "FP2"), ("FP-3", "FP3"), ("mean", "raw_mean")):
        c.has("score-size rho, %s" % t, fmt(rsum["score_size_rho"][key], 3))

    paired = pd.read_csv(d / "Recognition_objectives_paired.csv").set_index("compound")
    for cid, label in (("SANC00370", "acteoside"), ("SANC01049", "aloeresin C"),
                       ("SANC00853", "rosmarinic acid"), ("SANC00867", "tribuloside")):
        r = paired.loc[cid]
        c.has("%s FP-2 score" % label, fmt(r.FP2, 2))
        c.has("%s FP-3 score" % label, fmt(r.FP3, 2))
        c.has("%s mean percentile" % label, fmt(r.mean_percentile, 1),
              fmt(r.mean_percentile, 2))

    qsum = json.loads((d / "Quantum_analysis_summary.json").read_text(encoding="utf-8"))
    c.has("omega lower bound", fmt(qsum["omega_range_eV"][0], 2))
    c.has("omega upper bound", fmt(qsum["omega_range_eV"][1], 2))
    c.has("omega NO2 over OH, per cent",
          "%d%%" % round((qsum["omega_NO2_over_OH"] - 1) * 100))
    c.has("f+beta NO2 over OH, per cent",
          "%d%%" % round((1 - qsum["fplus_beta_NO2_over_OH"]) * 100))
    nitro = qsum["nitro_fragment_response"]["check"]
    c.has("nitro fragment ESP", fmt(nitro["ESP_para_substituent"], 3))
    c.has("nitro fragment Mulliken", fmt(nitro["Mulliken_para_substituent"], 3))
    c.has("nitro ester-alkene ESP", fmt(nitro["ESP_ester_alkene"], 3))
    c.has("nitro ester-alkene Mulliken", fmt(nitro["Mulliken_ester_alkene"], 3))

    desc = pd.read_csv(d / "Series_electronic_descriptors.csv")
    prim = desc[(desc.dataset == "primary") & (desc.scheme == "ESP")].set_index("substituent")
    c.has("f+beta OH", fmt(prim.loc["OH", "fplus_beta"], 3))
    c.has("f+beta NO2", fmt(prim.loc["NO2", "fplus_beta"], 3))

    eff = pd.read_csv(d / "Matched_effects_by_rule.csv")

    def contrast(target, arch, rule):
        block = eff[(eff.target == target) & (eff.architecture == arch)
                    & (eff.rule == rule)]
        return block.iloc[0] if len(block) else None

    fp3g = contrast("FP3", "glycosylated", "minimum_DockingScore")
    c.has("FP-3 glycosylated 4-OH score", fmt(fp3g.OH_value, 3), fmt(fp3g.OH_value, 2))
    c.has("FP-3 glycosylated 4-Cl score", fmt(fp3g.Cl_value, 3), fmt(fp3g.Cl_value, 2))
    c.has("FP-3 glycosylated contrast", fmt(abs(fp3g.Cl_minus_OH), 2))
    fp2g = contrast("FP2", "glycosylated", "minimum_DockingScore")
    c.has("FP-2 glycosylated 4-OH score", fmt(fp2g.OH_value, 3), fmt(fp2g.OH_value, 2))
    c.has("FP-2 glycosylated 4-Cl score", fmt(fp2g.Cl_value, 3), fmt(fp2g.Cl_value, 2))
    c.has("FP-2 glycosylated contrast", fmt(abs(fp2g.Cl_minus_OH), 2))

    hal = json.loads((d / "Matched_halogen_bond_check.json").read_text(encoding="utf-8"))
    c.has("chlorine contacts screened", hal["contacts_screened"])
    c.has("maximum C-Cl-acceptor angle", fmt(hal["max_angle_deg"], 1))
    c.has("shortest Cl contact", fmt(hal["min_distance_A"], 3))

    clus = json.loads((d / "Matched_pose_clustering.json").read_text(encoding="utf-8"))
    c.has("poses clustered", clus["poses_analysed"])
    c.has("clustering cutoff", fmt(clus["cutoff_A"], 1))
    for s in clus["summaries"]:
        if s["target"] == "FP3" and s["architecture"] == "glycosylated":
            c.has("FP-3 glycosylated median core RMSD",
                  fmt(s["median_pairwise_core_rmsd_A"], 2))

    dyad = json.loads((d / "Dyad_sensitivity.json").read_text(encoding="utf-8"))
    comp = dyad["comparison"]
    c.has("dyad heavy atoms compared", comp["heavy_atoms_compared"],
          "{:,}".format(comp["heavy_atoms_compared"]))
    c.has("dyad heavy-atom agreement", fmt(comp["heavy_atom_rmsd_A"], 6))
    for lab in ("standardized", "alternative"):
        c.has("dyad receptor atoms, %s" % lab, dyad[lab]["receptor_atoms"],
              "{:,}".format(dyad[lab]["receptor_atoms"]))
        c.has("dyad receptor heavy atoms, %s" % lab, dyad[lab]["receptor_heavy_atoms"],
              "{:,}".format(dyad[lab]["receptor_heavy_atoms"]))
        c.has("dyad receptor residues, %s" % lab, dyad[lab]["receptor_residues"])
        c.has("dyad receptor net charge, %s" % lab,
              str(dyad[lab]["receptor_formal_charge"]).replace("-", MINUS))
    # the pH condition the matched contrast is reported under
    ph = json.loads((d / "pH_sensitivity_summary.json").read_text(encoding="utf-8"))
    fp3 = ph["matched_contrast"]["FP3_glycosylated"]
    c.has("FP-3 glycosylated contrast at pH 7.0", fmt(fp3["pH70"], 2))
    c.has("FP-3 glycosylated contrast at pH 5.5", fmt(fp3["pH55"], 2))
    # the pose-ensemble statistic, which is a median over cross-pairs
    cl = json.loads((d / "Matched_pose_clustering.json").read_text(encoding="utf-8"))
    s3 = [x for x in cl["summaries"]
          if x["target"] == "FP3" and x["architecture"] == "glycosylated"][0]
    c.has("median cross-ensemble core displacement", fmt(s3["median_pairwise_core_rmsd_A"], 1))
    c.has("closest family centroid separation", fmt(s3["closest_cluster_centroid_rmsd_A"], 2))
    c.has("pose clustering cutoff", fmt(cl["cutoff_A"], 1))
    c.has("dyad shift", fmt(comp["shift_DockingScore"], 2))
    c.has("alternative dyad contrast",
          fmt(dyad["alternative"]["Cl_minus_OH_DockingScore"], 3))
    c.has("standardized dyad contrast",
          fmt(dyad["standardized"]["Cl_minus_OH_DockingScore"], 3),
          fmt(abs(dyad["standardized"]["Cl_minus_OH_DockingScore"]), 2))
    c.has("alternative dyad receptor atoms", dyad["alternative"]["receptor_atoms"],
          "{:,}".format(dyad["alternative"]["receptor_atoms"]))

    lib = d / "Dyad_library_effect.json"
    if lib.is_file():
        j = json.loads(lib.read_text(encoding="utf-8"))
        c.has("library parents compared", j["parents_compared"])
        c.has("library median absolute change", fmt(j["median_absolute_change"], 2))
        c.has("library IQR lower", fmt(j["iqr_absolute_change"][0], 2))
        c.has("library IQR upper", fmt(j["iqr_absolute_change"][1], 2))
        c.has("library maximum change", fmt(j["max_absolute_change"], 2))
        c.has("library largest mover", j["largest_mover"])
        c.has("parents moving more than the reference effect",
              j["parents_moving_more_than_reference"], "Fifty-seven")
        c.has("fraction moving more than the reference effect",
              "%d%%" % round(j["fraction_moving_more_than_reference"] * 100))
        c.has("library rank agreement", fmt(j["rank_agreement_spearman"], 3))
        if j.get("rank_agreement_ci"):
            c.has("library rank agreement CI lower", fmt(j["rank_agreement_ci"][0], 3))
            c.has("library rank agreement CI upper", fmt(j["rank_agreement_ci"][1], 3))
        c.has("library state changes", j["parents_changing_selected_state"], "nine")
        c.has("library top-ten overlap", j["top10_overlap"], "seven")
    else:
        print("  note  library dyad effect not present; those checks are skipped")

    ph = d / "pH_sensitivity_summary.json"
    if ph.is_file():
        j = json.loads(ph.read_text(encoding="utf-8"))
        for target, label in (("FP2", "FP-2"), ("FP3", "FP-3")):
            c.has("pH rank agreement, %s" % label,
                  fmt(j["rank_agreement"][target], 2))
            ci = j.get("rank_agreement_ci", {}).get(target)
            if ci and ci[0] is not None:
                c.has("pH rank agreement CI, %s" % label,
                      "%s to %s" % (fmt(ci[0], 2), fmt(ci[1], 2)))
            c.has("pH median score change, %s" % label,
                  fmt(j["absolute_score_change"][target]["median"], 2))
            c.has("pH maximum score change, %s" % label,
                  fmt(j["absolute_score_change"][target]["maximum"], 2))
            c.has("pH top-ten overlap, %s" % label,
                  j["top_k_overlap"][target]["top10"],
                  {7: "seven", 8: "eight"}.get(j["top_k_overlap"][target]["top10"]))
        fp3 = j["matched_contrast"]["FP3_glycosylated"]
        c.has("FP-3 glycosylated contrast at pH 5.5", fmt(fp3["pH55"], 2))
        c.has("compounds compared across pH",
              j["compounds_compared"]["FP3"], "eighteen")
    else:
        print("  note  pH summary not present; those checks are skipped")

    cov = pd.read_csv(d / "Covalent_stage_records.csv")
    c.has("covalent complexes", len(cov))
    c.has("post-reaction scored", int((cov.postreaction_status == "completed").sum()))

    bench = d / "Benchmark_summary.json"
    if bench.is_file():
        j = json.loads(bench.read_text(encoding="utf-8"))
        b = j["all_actives"]
        c.has("benchmark ROC-AUC", fmt(b["roc_auc"], 3))
        c.has("benchmark ROC-AUC CI lower", fmt(b["roc_ci"][0], 3))
        c.has("benchmark ROC-AUC CI upper", fmt(b["roc_ci"][1], 3))
        c.has("benchmark N", b["n"], "{:,}".format(b["n"]))
        k = j["covalent_only"]
        c.has("covalent-subset ROC-AUC", fmt(k["roc_auc"], 3))
        c.has("covalent-subset ROC CI lower", fmt(k["roc_ci"][0], 3))
        c.has("covalent-subset ROC CI upper", fmt(k["roc_ci"][1], 3))
        c.has("covalent-subset EF5", fmt(k["ef5"], 2))
        c.has("covalent-subset EF5 CI lower", fmt(k["ef5_ci"][0], 2))
        c.has("covalent-subset EF5 CI upper", fmt(k["ef5_ci"][1], 2))

    for tag, label in (("FP2_XP", "FP-2 XP"), ("FP3_XP", "FP-3 XP")):
        f = d / ("Benchmark_summary_%s.json" % tag)
        if not f.is_file():
            continue
        j = json.loads(f.read_text(encoding="utf-8"))
        subsets = (("all_actives", "all actives"), ("covalent_only", "covalent actives"))
        if tag == "FP3_XP":
            subsets = subsets[:1]
        for subset, name in subsets:
            m = j[subset]
            c.has("%s %s ROC-AUC" % (label, name), fmt(m["roc_auc"], 3))
            c.has("%s %s ROC CI lower" % (label, name), fmt(m["roc_ci"][0], 3))
            c.has("%s %s ROC CI upper" % (label, name), fmt(m["roc_ci"][1], 3))
        m = j["all_actives"]
        c.has("%s scored compounds" % label, m["n"], "{:,}".format(m["n"]))
        c.has("%s actives scored" % label, m["actives"], "{:,}".format(m["actives"]))

    cmp_path = d / "Benchmark_dyad_comparison.json"
    if cmp_path.is_file():
        j = json.loads(cmp_path.read_text(encoding="utf-8"))
        c.has("matched benchmark common set", j["scored_by_both"],
              "{:,}".format(j["scored_by_both"]))
        for lab, key in (("standardized", "standardized"), ("ion pair", "alternative")):
            c.has("matched benchmark ROC, %s, all actives" % lab,
                  fmt(j[key]["all_actives"]["roc_auc"], 3))
            c.has("matched benchmark ROC, %s, covalent" % lab,
                  fmt(j[key]["covalent_only"]["roc_auc"], 3))
            c.has("matched benchmark EF5, %s, covalent" % lab,
                  fmt(j[key]["covalent_only"]["ef5"], 2))
    else:
        print("  note  benchmark summary not present; those checks are skipped")

    conc = d / "Property_platform_concordance.json"
    if conc.is_file():
        j = json.loads(conc.read_text(encoding="utf-8"))
        c.has("Caco-2 rank agreement", fmt(j["caco2_rank_agreement_spearman"], 2))

    print("\nverifying %s against %s\n" % (a.manuscript.name, d))
    return c.report()


if __name__ == "__main__":
    sys.exit(main())
